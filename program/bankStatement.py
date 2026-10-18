#!/usr/bin/env python3
"""
bankStatement.py

Reads a plain text / notepad file containing multiple stacked CSV blocks
(one "Checks" table and one "Deposits" table per month, repeated for as
many months as you like) and produces a single Excel workbook with two
sheets:

    1. Checks and Withdrawals
    2. Deposits and Sales

Rows are sorted by Date within each sheet. All months found in the input
file are combined into one workbook.

USAGE
-----
    python bankStatement.py input.txt output.xlsx

Requires:
    pip install pandas openpyxl
"""

import sys
import csv
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


CHECKS_HEADER = re.compile(
    r"^\s*date\s*,\s*payee",
    re.IGNORECASE
)

DEPOSITS_HEADER = re.compile(
    r"^\s*date\s*,\s*received\s+from",
    re.IGNORECASE
)


def split_blocks(text: str):
    lines = text.splitlines()

    blocks = []
    current_type = None
    current_lines = []

    for raw_line in lines:
        line = raw_line.rstrip("\n")

        is_checks_header = CHECKS_HEADER.match(line) is not None
        is_deposits_header = DEPOSITS_HEADER.match(line) is not None

        if is_checks_header or is_deposits_header:
            if current_type is not None and current_lines:
                blocks.append((current_type, current_lines))

            current_type = "checks" if is_checks_header else "deposits"
            current_lines = [line]
            continue

        if current_type is not None:
            current_lines.append(line)

    if current_type is not None and current_lines:
        blocks.append((current_type, current_lines))

    return blocks


def parse_block_to_dataframe(block_lines):
    if not block_lines:
        return pd.DataFrame()

    lines = [line for line in block_lines if line.strip()]

    if not lines:
        return pd.DataFrame()

    reader = csv.reader(lines)

    try:
        rows = list(reader)
    except csv.Error as e:
        print(f"Warning: CSV parsing error: {e}")
        return pd.DataFrame()

    if not rows:
        return pd.DataFrame()

    header = [value.strip() for value in rows[0]]
    expected_cols = len(header)

    valid_rows = []
    tail_count = 3

    for row in rows[1:]:
        if not any(value.strip() for value in row):
            continue

        if len(row) == expected_cols:
            valid_rows.append(row)
            continue

        if len(row) > expected_cols:
            if len(row) <= tail_count:
                print("Warning: skipped malformed row:", row)
                continue

            date_field = row[0]
            tail_fields = row[-tail_count:]
            middle_fields = row[1:-tail_count]
            middle_target = expected_cols - 1 - tail_count

            if middle_target <= 0:
                print("Warning: skipped malformed row:", row)
                continue

            if middle_target == 1:
                merged_middle = [",".join(middle_fields)]
            else:
                merged_middle = (
                    middle_fields[:middle_target - 1]
                    + [",".join(middle_fields[middle_target - 1:])]
                )

            fixed_row = [date_field] + merged_middle + tail_fields

            if len(fixed_row) == expected_cols:
                valid_rows.append(fixed_row)
            else:
                print("Warning: could not repair row:", row)

            continue

        print("Warning: skipped row with too few columns:", row)

    if not valid_rows:
        return pd.DataFrame(columns=header)

    return pd.DataFrame(valid_rows, columns=header)


def combine_blocks(blocks, block_type):
    frames = []

    for btype, lines in blocks:
        if btype != block_type:
            continue

        df = parse_block_to_dataframe(lines)

        if not df.empty:
            frames.append(df)

    if not frames:
        return pd.DataFrame()

    return pd.concat(frames, ignore_index=True)


def sort_by_date(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or "Date" not in df.columns:
        return df

    parsed = pd.to_datetime(df["Date"], format="%m/%d/%Y", errors="coerce")

    df = df.assign(_sort_key=parsed)
    df = df.sort_values(by=["_sort_key"], kind="stable", na_position="last")
    df = df.drop(columns=["_sort_key"])
    df = df.reset_index(drop=True)

    return df


def autosize_and_style(writer, sheet_name, df):
    ws = writer.sheets[sheet_name]

    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        header_cell = ws[f"{col_letter}1"]

        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal="center")

        max_len = len(str(col_name))

        for value in df[col_name]:
            max_len = max(max_len, len(str(value)))

        ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = body_font

    if "Amount" in df.columns:
        amount_col_idx = list(df.columns).index("Amount") + 1
        amount_letter = get_column_letter(amount_col_idx)

        for row_num in range(2, ws.max_row + 1):
            cell = ws[f"{amount_letter}{row_num}"]
            cell.number_format = "#,##0.00"

    ws.freeze_panes = "A2"


def main():
    if len(sys.argv) != 3:
        print("Usage: python bankStatement.py <input.txt> <output.xlsx>")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    if not input_path.exists():
        print(f"Input file not found: {input_path}")
        sys.exit(1)

    text = input_path.read_text(encoding="utf-8", errors="replace")

    blocks = split_blocks(text)

    if not blocks:
        print("No 'Date,Payee...' or 'Date,Received From...' header rows were found.")
        print("Check that your CSV blocks start with those exact header rows.")
        sys.exit(1)

    checks_df = combine_blocks(blocks, "checks")
    deposits_df = combine_blocks(blocks, "deposits")

    checks_df = sort_by_date(checks_df)
    deposits_df = sort_by_date(deposits_df)

    if checks_df.empty and deposits_df.empty:
        print("Found header rows but could not parse any CSV data rows out of them.")
        sys.exit(1)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if not checks_df.empty:
            checks_df.to_excel(writer, sheet_name="Checks and Withdrawals", index=False)
            autosize_and_style(writer, "Checks and Withdrawals", checks_df)
        else:
            print("Warning: no Checks and Withdrawals rows found — sheet not created.")

        if not deposits_df.empty:
            deposits_df.to_excel(writer, sheet_name="Deposits and Sales", index=False)
            autosize_and_style(writer, "Deposits and Sales", deposits_df)
        else:
            print("Warning: no Deposits and Sales rows found — sheet not created.")

    print(
        f"Done. Wrote {len(checks_df)} check/withdrawal rows and "
        f"{len(deposits_df)} deposit/sales rows to: {output_path}"
    )


if __name__ == "__main__":
    main()